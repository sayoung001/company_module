"""읽기 전용 접근 (SPEC §1.1, §1.2, §20.1).

* ``.xlsx`` → openpyxl (``read_only=True``)
* ``.xls``  → xlrd 2.x
* **어느 쪽도 파일을 수정하지 않는다.** 엑셀도 띄우지 않는다.

배지 계산·판정·사후 검증처럼 값만 필요한 경로는 전부 여기를 쓴다.
xlwings 는 쓰기 전용이다.
"""
from __future__ import annotations

from contextlib import contextmanager
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterator

from .sheets import col_index


class ReaderError(RuntimeError):
    pass


@dataclass
class ReadSheet:
    """읽기 전용 시트. 좌표는 1-based."""

    name: str
    _rows: list[list[Any]]

    @property
    def nrows(self) -> int:
        return len(self._rows)

    @property
    def ncols(self) -> int:
        return max((len(r) for r in self._rows), default=0)

    def cell(self, row: int, col: int) -> Any:
        if 1 <= row <= len(self._rows):
            r = self._rows[row - 1]
            if 1 <= col <= len(r):
                return r[col - 1]
        return None

    def at(self, ref: str) -> Any:
        """at('DG5') — A1 표기로 읽기."""
        i = 0
        while i < len(ref) and ref[i].isalpha():
            i += 1
        return self.cell(int(ref[i:]), col_index(ref[:i]))

    def row(self, row: int) -> list[Any]:
        return list(self._rows[row - 1]) if 1 <= row <= len(self._rows) else []

    def col_values(self, col: int, start: int = 1, end: int | None = None) -> list[Any]:
        end = end or self.nrows
        return [self.cell(r, col) for r in range(start, end + 1)]

    def last_filled_row(self, col: int, start: int = 1) -> int:
        last = start - 1
        for r in range(start, self.nrows + 1):
            v = self.cell(r, col)
            if v is not None and str(v).strip() != "":
                last = r
        return last

    def iter_rows(self, start: int = 1, end: int | None = None) -> Iterator[tuple[int, list[Any]]]:
        for r in range(start, (end or self.nrows) + 1):
            yield r, self.row(r)


@dataclass
class ReadBook:
    path: Path
    sheets: dict[str, ReadSheet]

    @property
    def sheet_names(self) -> list[str]:
        return list(self.sheets)

    def sheet(self, name: str) -> ReadSheet:
        if name not in self.sheets:
            raise ReaderError(
                f"시트가 없습니다: {name!r} in {self.path.name} "
                f"(있는 시트: {self.sheet_names})"
            )
        return self.sheets[name]

    def find_sheet(self, *candidates: str) -> ReadSheet | None:
        for c in candidates:
            if c in self.sheets:
                return self.sheets[c]
        return None


# ---------------------------------------------------------------------
def read_book(path: str | Path, *, data_only: bool = True,
              only: list[str] | None = None) -> ReadBook:
    """확장자에 맞는 엔진으로 통합문서를 통째로 읽는다.

    ``data_only=True`` 면 수식 대신 마지막 계산값을 읽는다 (사후 검증용).
    ``only`` 로 시트를 좁히면 큰 파일에서 훨씬 빠르다.
    """
    p = Path(path)
    if not p.exists():
        raise ReaderError(f"파일이 없습니다: {p}")
    ext = p.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        return _read_xlsx(p, data_only=data_only, only=only)
    if ext == ".xls":
        return _read_xls(p, only=only)
    raise ReaderError(f"읽을 수 없는 확장자입니다: {p.name}")


def _read_xlsx(p: Path, *, data_only: bool, only: list[str] | None) -> ReadBook:
    try:
        import openpyxl
    except ImportError as e:                      # pragma: no cover
        raise ReaderError("openpyxl 이 필요합니다: pip install openpyxl") from e
    # read_only=True 는 파일을 수정하지 않는다. 경고(wmf 이미지 드롭)는 저장할 때만
    # 문제가 되는데, 여기서는 저장 자체를 하지 않는다.
    wb = openpyxl.load_workbook(p, data_only=data_only, read_only=True, keep_links=False)
    try:
        sheets: dict[str, ReadSheet] = {}
        for ws in wb.worksheets:
            if only and ws.title not in only:
                continue
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            sheets[ws.title] = ReadSheet(ws.title, rows)
        return ReadBook(p, sheets)
    finally:
        wb.close()


def _read_xls(p: Path, *, only: list[str] | None) -> ReadBook:
    try:
        import xlrd
    except ImportError as e:                      # pragma: no cover
        raise ReaderError(
            "레거시 .xls 를 읽으려면 xlrd 2.x 가 필요합니다: pip install xlrd==2.0.1"
        ) from e
    book = xlrd.open_workbook(str(p), formatting_info=False, on_demand=True)
    try:
        sheets: dict[str, ReadSheet] = {}
        for name in book.sheet_names():
            if only and name not in only:
                continue
            sh = book.sheet_by_name(name)
            rows: list[list[Any]] = []
            for r in range(sh.nrows):
                out: list[Any] = []
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    out.append(_xls_value(cell, book.datemode))
                rows.append(out)
            sheets[name] = ReadSheet(name, rows)
        return ReadBook(p, sheets)
    finally:
        book.release_resources()


def _xls_value(cell: Any, datemode: int) -> Any:
    import xlrd

    if cell.ctype == xlrd.XL_CELL_DATE:
        y, mo, d, h, mi, s = xlrd.xldate_as_tuple(cell.value, datemode)
        from datetime import date, datetime

        return date(y, mo, d) if (h, mi, s) == (0, 0, 0) else datetime(y, mo, d, h, mi, s)
    if cell.ctype == xlrd.XL_CELL_EMPTY:
        return None
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    if cell.ctype == xlrd.XL_CELL_ERROR:
        return None
    return cell.value


@contextmanager
def open_read(path: str | Path, **kw: Any) -> Iterator[ReadBook]:
    yield read_book(path, **kw)
