"""백업·복원, 미리보기, 현황 덤프 (SPEC §2.1, §2.2, §7, §12)."""
from __future__ import annotations

from datetime import date, datetime

import openpyxl
import pytest

from core.backup import BackupError, BackupManager
from core.config import Config
from core.dump import dump_state
from core.sheets import col_index, col_letter
from core.util import is_locked, sha256
from tests.fixtures import cfg as base_cfg


# =====================================================================
# §2.1 백업
# =====================================================================
def test_백업후_복원하면_해시가_같다(tmp_path):
    """S2 완료 기준 — 백업 후 복원하면 파일 해시가 동일."""
    src = tmp_path / "원본.xlsx"
    openpyxl.Workbook().save(src)
    before = sha256(src)

    m = BackupManager(tmp_path / "백업")
    bset = m.start("자재검수_회차생성")
    m.add(bset, src)
    m.write_manifest(bset, {"charge_no": 13})

    src.write_text("망가뜨림", encoding="utf-8")
    assert sha256(src) != before
    bset.restore()
    assert sha256(src) == before


def test_manifest_에_입력값과_해시가_남는다(tmp_path):
    import json

    src = tmp_path / "원본.xlsx"
    openpyxl.Workbook().save(src)
    m = BackupManager(tmp_path / "백업")
    bset = m.start("자재검수")
    m.add(bset, src)
    path = m.write_manifest(bset, {"charge_no": 13, "input": {"차수": 13}})
    data = json.loads(path.read_text(encoding="utf-8"))
    assert data["task"] == "자재검수"
    assert data["files"] == ["원본.xlsx"]
    assert data["charge_no"] == 13
    assert data["sha256"]["원본.xlsx"] == sha256(src)


def test_원본이_없으면_백업은_실패한다(tmp_path):
    m = BackupManager(tmp_path / "백업")
    bset = m.start("t")
    with pytest.raises(BackupError):
        m.add(bset, tmp_path / "없는파일.xlsx")


def test_같은날_여러번_실행해도_덮어쓰지_않는다(tmp_path):
    src = tmp_path / "a.xlsx"
    openpyxl.Workbook().save(src)
    m = BackupManager(tmp_path / "백업")
    폴더 = set()
    for sec in (1, 2):
        bset = m.start("t", when=datetime(2026, 9, 5, 14, 3, sec))
        m.add(bset, src)
        m.write_manifest(bset, {})
        폴더.add(bset.folder)
    assert len(폴더) == 2


def test_백업목록은_최신순(tmp_path):
    src = tmp_path / "a.xlsx"
    openpyxl.Workbook().save(src)
    m = BackupManager(tmp_path / "백업")
    for sec in (1, 5, 9):
        bset = m.start("t", when=datetime(2026, 9, 5, 14, 3, sec))
        m.add(bset, src)
        m.write_manifest(bset, {})
    목록 = m.list_backups()
    assert len(목록) == 3
    assert [b.created.second for b in 목록] == [9, 5, 1]


# =====================================================================
# §10 #7 파일 잠금
# =====================================================================
def test_엑셀_잠금파일이_있으면_잠긴_것으로_본다(tmp_path):
    f = tmp_path / "갑지.xlsx"
    openpyxl.Workbook().save(f)
    assert is_locked(f) is False
    (tmp_path / "~$갑지.xlsx").write_bytes(b"")
    assert is_locked(f) is True


# =====================================================================
# 열 문자 변환 — 실측 블록 열
# =====================================================================
@pytest.mark.parametrize("col,letter", [
    (1, "A"), (11, "K"), (17, "Q"), (33, "AG"), (61, "BI"), (73, "BU"),
    (78, "BZ"), (89, "CK"), (101, "CW"), (111, "DG"), (120, "DP"),
])
def test_열문자_변환(col, letter):
    assert col_letter(col) == letter
    assert col_index(letter) == col


# =====================================================================
# §12 현황 덤프
# =====================================================================
def test_덤프가_실측값을_그대로_출력한다(tmp_path, capsys):
    체크 = tmp_path / "체크용.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "파일 (2)"
    for i in range(37):                          # 18~54행
        r = 18 + i
        ws.cell(r, 1, date(2026, 8, 6))
        ws.cell(r, 7, 6)
    ws.cell(54, 11, 12)                          # 최신 차수 12
    wb.save(체크)

    갑지 = tmp_path / "갑지.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "자재검수(PHC)"
    for n in range(1, 12):                       # 블록 11개 = 차수 12까지
        c0 = 1 + (n - 1) * 10
        ws.cell(2, c0 + 2, f"Q-G-03(파일)-{n + 1:02d}")
        ws.cell(9, c0 + 7, 222)
    wb.save(갑지)

    수불부 = tmp_path / "수불부.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "500"
    ws.cell(1, 1, "주요자재 검사 및 수불부 (26년 08월)")
    ws.cell(1, 17, "주요자재 검사 및 수불부 (26년 09월)")
    for r in (5, 7, 9):
        ws.cell(r, 17 + 4, date(2026, 9, 4))
        ws.cell(r, 17 + 7, 222)
    wb.save(수불부)

    c = base_cfg()
    raw = dict(c.raw)
    raw["files"] = dict(raw["files"])
    raw["files"].update(supply_check=str(체크), request_cover=str(갑지),
                        ledger_phc=str(수불부))
    dump_state(Config(raw=raw))

    out = capsys.readouterr().out
    assert "마지막 데이터 행 : 54" in out
    assert "최신 차수        : 12  ->  다음 회차 13" in out
    assert "마지막 블록      : 11번  시작 열 CW (101)" in out
    assert "다음 회차 블록   : DG (111)" in out
    assert "26년 09월" in out
    assert "마지막 데이터 행 : 9" in out
