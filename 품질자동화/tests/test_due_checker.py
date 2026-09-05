"""시험 도래 판정 검증 (SPEC §16.0, §17.0, §18.2).

§16.0 의 '2026-09-05 실측 현황' 표가 정답지다. 실제 파일과 같은 구조의
xlsx 를 만들어 그 표가 그대로 재현되는지 본다. 주입 시나리오 4종도 함께.
"""
from __future__ import annotations

from datetime import date

import openpyxl
import pytest

from core.config import Config
from core.due_checker import DueChecker
from core.state import PendingTest, State
from core.util import sha256
from tests.fixtures import cfg as base_cfg

# §16.0 실측 현황 — 마지막 시험일과 그 이후 누적
실측 = {
    "KCC":          (date(2026, 8, 20), 69),
    "아주산업":     (date(2026, 8, 20), 24),
    "한국파일":     (date(2026, 8, 24), 24),
    "동양파일":     (date(2026, 8, 28), 21),
    "삼일씨엔에스": (date(2026, 9, 3), 0),
    "동진파일":     (date(2026, 9, 4), 0),
}
시험전_반입 = 14        # 각 업체가 시험일 이전에 받은 양 (합계 222본을 맞춘다)


def _체크용(path, 추가: list[tuple[str, date, int]] | None = None) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "파일 (2)"
    ws["A17"], ws["C17"], ws["G17"] = "송장일자", "업체명", "수량\n(본)"
    r = 18
    for 업체, (시험일, 누적) in 실측.items():
        ws.cell(r, 1, 시험일)                      # 시험일 당일 반입 (누적에 안 들어감)
        ws.cell(r, 3, 업체)
        ws.cell(r, 7, 시험전_반입)
        r += 1
        if 누적:
            ws.cell(r, 1, date(2026, 9, 5))        # 시험 이후 반입
            ws.cell(r, 3, 업체)
            ws.cell(r, 7, 누적)
            r += 1
    for 업체, 일자, 본 in (추가 or []):
        ws.cell(r, 1, 일자)
        ws.cell(r, 3, 업체)
        ws.cell(r, 7, 본)
        r += 1
    wb.save(path)


def _601대장(path, 생략: set[str] | None = None) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "대장"
    ws["A4"], ws["B4"], ws["C4"] = "No", "날짜", "업체명"
    r = 5
    표기 = {"KCC": "㈜KCC글라스", "아주산업": "아주산업㈜", "한국파일": "한국파일㈜",
            "동양파일": "동양파일㈜", "삼일씨엔에스": "㈜삼일C&S",
            "동진파일": "동진파일㈜ "}          # 뒤 공백 — 실제 파일에 존재한다
    for i, (업체, (시험일, _)) in enumerate(실측.items(), start=1):
        if 생략 and 업체 in 생략:
            continue
        ws.cell(r, 1, i)
        ws.cell(r, 2, 시험일)
        ws.cell(r, 3, 표기[업체])
        r += 1
    wb.save(path)


def _612대장(path, 날짜들: list[date]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "대장"
    for i, d in enumerate(날짜들):
        ws.cell(5 + i, 1, i + 1)
        ws.cell(5 + i, 2, d)
    wb.save(path)


@pytest.fixture()
def 환경(tmp_path):
    """실제 파일과 같은 구조의 임시 파일 3개 + 그걸 가리키는 config."""
    def build(추가=None, 생략=None, 밀크=None):
        c = base_cfg()
        raw = dict(c.raw)
        raw["files"] = dict(raw["files"])
        raw["paths"] = dict(raw["paths"])
        raw["paths"]["state_dir"] = str(tmp_path / "상태")
        for key, name, fn in (
            ("supply_check", "체크용.xlsx", lambda p: _체크용(p, 추가)),
            ("test_shape", "601.xlsx", lambda p: _601대장(p, 생략)),
            ("test_milk", "612.xlsx", lambda p: _612대장(p, 밀크 or [])),
        ):
            p = tmp_path / name
            fn(p)
            raw["files"][key] = str(p)
        return Config(raw=raw)
    return build


# =====================================================================
def test_실측표_재현_아무도_도래하지_않는다(환경):
    """§16.0 — 6개 업체 모두 이미 1회씩 시험을 마쳤고 200본에 걸린 곳이 없다."""
    alerts = DueChecker(환경()).check(date(2026, 9, 5))
    현황 = {v.업체: v for v in alerts.all_vendors}

    assert len(현황) == 6
    for 업체, (시험일, 누적) in 실측.items():
        assert 현황[업체].마지막시험 == 시험일, 업체
        assert 현황[업체].누적 == 누적, 업체
        assert 현황[업체].level == "ok", 업체
    assert alerts.due_count == 0
    assert alerts.warnings == []
    # 전체 반입 222본
    assert sum(v.누적 for v in alerts.all_vendors) + 시험전_반입 * 6 == 222


def test_주입_KCC_140본이면_209본_due(환경):
    """§16.0 검증 — KCC 69 + 140 = 209본 > 200 -> 🔴"""
    alerts = DueChecker(환경(추가=[("KCC", date(2026, 9, 5), 140)])).check(date(2026, 9, 5))
    kcc = next(v for v in alerts.all_vendors if v.업체 == "KCC")
    assert kcc.level == "due"
    assert kcc.누적 == 209
    assert "209본" in kcc.사유 and "200본" in kcc.사유
    assert alerts.due_count == 1


def test_주입_시험이력없는_업체는_최초반입_due(환경):
    alerts = DueChecker(환경(추가=[("아이에스동서", date(2026, 9, 5), 12)])).check(date(2026, 9, 5))
    v = next(v for v in alerts.all_vendors if v.업체 == "아이에스동서")
    assert v.level == "due"
    assert "최초 반입" in v.사유 and "12본" in v.사유


def test_주입_181본은_warn_20본_남음(환경):
    """181 + 20 = 201 > 200 이므로 '20본 남음'."""
    alerts = DueChecker(환경(추가=[("동양파일", date(2026, 9, 5), 160)])).check(date(2026, 9, 5))
    v = next(v for v in alerts.all_vendors if v.업체 == "동양파일")
    assert v.누적 == 181
    assert v.level == "warn"
    assert v.남은본 == 20
    assert "20본 남음" in str(v)


def test_모르는_업체명은_판정에서_빼고_반드시_알린다(환경):
    """§16.0 — 조용히 무시하면 그 업체는 영원히 누락된다. 가장 위험한 실패 모드."""
    alerts = DueChecker(환경(추가=[("없는회사파일", date(2026, 9, 5), 30)])).check(date(2026, 9, 5))
    assert not any(v.업체 == "없는회사파일" for v in alerts.all_vendors)
    assert any("없는회사파일" in w and "vendor_alias" in w for w in alerts.warnings)


def test_check는_원본을_수정하지_않는다(환경):
    cfg = 환경()
    before = {k: sha256(cfg.path(k)) for k in ("supply_check", "test_shape", "test_milk")}
    DueChecker(cfg).check(date(2026, 9, 5))
    after = {k: sha256(cfg.path(k)) for k in ("supply_check", "test_shape", "test_milk")}
    assert before == after


def test_alerts_json_이_남는다(환경, tmp_path):
    import json

    DueChecker(환경()).check(date(2026, 9, 5))
    data = json.loads((tmp_path / "상태" / "alerts.json").read_text(encoding="utf-8"))
    assert data["due_count"] == 0
    assert len(data["shape"]) == 6


# =====================================================================
# §17.0 밀크 주간 카운터
# =====================================================================
def test_밀크_주간카운터_2회면_1회_남음(환경):
    """2026-09-03(목) 기준. 그 주 월요일은 08-31, 남은 평일은 목·금 2일."""
    alerts = DueChecker(환경(밀크=[date(2026, 8, 31), date(2026, 9, 2)])).check(date(2026, 9, 3))
    w = alerts.milk
    assert (w.실시, w.목표) == (2, 3)
    assert w.남은평일 == 2                # 목(9/3) · 금(9/4)
    assert w.남은횟수 == 1
    assert w.위험 is False
    assert "●●○" in str(w)


def test_밀크_남은평일보다_남은횟수가_많으면_빨간색(환경):
    """§17.0 예시 — 금요일에 1회뿐이면 2회 남았는데 평일은 1일뿐 -> 🔴"""
    alerts = DueChecker(환경(밀크=[date(2026, 8, 31)])).check(date(2026, 9, 4))
    w = alerts.milk                       # 9/4 는 금요일
    assert (w.실시, w.남은평일, w.남은횟수) == (1, 1, 2)
    assert w.위험 is True
    assert str(w).startswith("🔴")
    # 토요일이 되면 평일이 하나도 안 남는다
    w2 = DueChecker.밀크_주간현황([date(2026, 8, 31)], date(2026, 9, 5))
    assert w2.남은평일 == 0 and w2.위험 is True


def test_밀크는_지난주_실시를_세지_않는다(환경):
    alerts = DueChecker(환경(밀크=[date(2026, 8, 25), date(2026, 8, 26)])).check(date(2026, 9, 4))
    assert alerts.milk.실시 == 0


# =====================================================================
# §18.2 진행 중 시험 추적
# =====================================================================
def test_진행중_시험_예정일_초과는_빨간색(환경, tmp_path):
    st = State(tmp_path / "state.json")
    st.upsert_pending(PendingTest(series="BSCW", 번호=7, 채취일="2026-08-25",
                                  목록행번호=7, stage="AWAIT_7D"))
    st.upsert_pending(PendingTest(series="BSCW", 번호=9, 채취일="2026-09-04",
                                  목록행번호=9, stage="AWAIT_7D"))
    alerts = DueChecker(환경(), state=st).check(date(2026, 9, 5))

    지연 = next(p for p in alerts.pending if p.시험번호 == "Q-Q-03-07")
    assert 지연.예정일 == date(2026, 9, 1)
    assert 지연.지연일수 == 4 and 지연.지연 is True
    assert "지연" in str(지연)

    대기 = next(p for p in alerts.pending if p.시험번호 == "Q-Q-03-09")
    assert 대기.예정일 == date(2026, 9, 11) and 대기.지연 is False
    assert alerts.due_count == 1                # 지연 1건


def test_state_는_재시작해도_유지된다(tmp_path):
    p = tmp_path / "state.json"
    st = State(p)
    st.set_check("자재검수", 13, "photo_paste", True)
    st.upsert_pending(PendingTest(series="JSP", 번호=1, 채취일="2026-08-27", 목록행번호=1))
    again = State(p)
    assert again.get_checks("자재검수", 13)["photo_paste"] is True
    assert again.pending[0].시험번호 == "Q-Q-04-01"
    assert again.open_workflows() == []          # 체크된 것뿐이면 미완료 없음
