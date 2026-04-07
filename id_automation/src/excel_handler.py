"""
Excel 파일 처리 핸들러
종합양식 ↔ 출석부양식 데이터 복사
"""

from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
import os
from datetime import datetime


class ExcelHandler:
    def __init__(self):
        # 컬럼 매핑 (0-indexed)
        self.columns = {
            'no': 0,           # A: No.
            'name': 1,         # B: 성명
            'birth': 2,        # C: 생년월일
            'phone': 3,        # D: 전화번호
            'nation': 4,       # E: 국적
            'support_type': 5, # F: 지원구분
            'support_target': 6, # G: 지원대상구분
            'fee_type': 7,     # H: 교육비부담여부
            'cert_type': 8,    # I: 이수증유형
            'payment': 9,      # J: 결제 정보
            'vulnerable': 10   # K: 취약 구분
        }
    
    def read_source_data(self, source_path: str) -> list:
        """종합양식에서 출석부 데이터 읽기"""
        wb = load_workbook(source_path, data_only=True)
        
        # '출석부양식' 시트 찾기
        sheet_name = None
        for name in wb.sheetnames:
            if '출석부' in name:
                sheet_name = name
                break
        
        if not sheet_name:
            raise ValueError("'출석부양식' 시트를 찾을 수 없습니다.")
        
        ws = wb[sheet_name]
        data = []
        
        # 2행부터 데이터 읽기 (1행은 헤더)
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, max_col=11), start=2):
            # 이름이 없으면 스킵
            name = row[self.columns['name']].value
            if not name or str(name).strip() == '' or name == 'NaN':
                continue
            
            # 생년월일 처리
            birth = row[self.columns['birth']].value
            if birth:
                if isinstance(birth, datetime):
                    birth = birth.strftime('%Y%m%d')
                else:
                    birth = str(birth).replace('-', '').replace('.', '')[:8]
            
            # 전화번호 처리
            phone = row[self.columns['phone']].value
            if phone:
                phone = str(phone).replace('-', '')
            
            row_data = {
                'no': row[self.columns['no']].value,
                'name': name,
                'birth': birth,
                'phone': phone,
                'nation': row[self.columns['nation']].value or 410,
                'support_type': row[self.columns['support_type']].value or '00',
                'support_target': row[self.columns['support_target']].value or '00',
                'fee_type': row[self.columns['fee_type']].value or 'SS',
                'cert_type': row[self.columns['cert_type']].value or 'MB',
                'payment': row[self.columns['payment']].value or '',
                'vulnerable': row[self.columns['vulnerable']].value or '0'
            }
            data.append(row_data)
        
        wb.close()
        return data
    
    def copy_to_attendance(self, source_path: str, target_path: str) -> dict:
        """종합양식 데이터를 출석부양식으로 복사"""
        try:
            # 소스 데이터 읽기
            data = self.read_source_data(source_path)
            
            if not data:
                return {
                    'success': False,
                    'message': '복사할 데이터가 없습니다.',
                    'count': 0
                }
            
            # 타겟 파일 열기
            wb = load_workbook(target_path)
            
            # '출석부양식' 시트 찾기
            sheet_name = None
            for name in wb.sheetnames:
                if '출석부' in name:
                    sheet_name = name
                    break
            
            if not sheet_name:
                wb.close()
                return {
                    'success': False,
                    'message': "'출석부양식' 시트를 찾을 수 없습니다.",
                    'count': 0
                }
            
            ws = wb[sheet_name]
            
            # 스타일 정의
            thin_border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
            center_align = Alignment(horizontal='center', vertical='center')
            
            # 데이터 쓰기 (2행부터)
            for idx, row_data in enumerate(data, start=2):
                ws.cell(row=idx, column=1, value=idx-1)  # No.
                ws.cell(row=idx, column=2, value=row_data['name'])
                ws.cell(row=idx, column=3, value=row_data['birth'])
                ws.cell(row=idx, column=4, value=row_data['phone'])
                ws.cell(row=idx, column=5, value=row_data['nation'])
                ws.cell(row=idx, column=6, value=row_data['support_type'])
                ws.cell(row=idx, column=7, value=row_data['support_target'])
                ws.cell(row=idx, column=8, value=row_data['fee_type'])
                ws.cell(row=idx, column=9, value=row_data['cert_type'])
                ws.cell(row=idx, column=10, value=row_data['payment'])
                ws.cell(row=idx, column=11, value=row_data['vulnerable'])
                
                # 스타일 적용
                for col in range(1, 12):
                    cell = ws.cell(row=idx, column=col)
                    cell.border = thin_border
                    cell.alignment = center_align
            
            # 저장 (원본 보존을 위해 새 파일로)
            base_name = os.path.splitext(os.path.basename(target_path))[0]
            dir_name = os.path.dirname(target_path)
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            new_path = os.path.join(dir_name, f"{base_name}_{timestamp}.xlsx")
            
            wb.save(new_path)
            wb.close()
            
            return {
                'success': True,
                'message': '복사 완료',
                'count': len(data),
                'saved_path': new_path
            }
            
        except Exception as e:
            return {
                'success': False,
                'message': str(e),
                'count': 0
            }
    
    def save_attendance(self, source_path: str, output_dir: str = "",
                        date_str: str = "") -> dict:
        """
        종합양식의 출석부양식 시트에서 A~I열, 헤더+N명 데이터를
        출석부양식_YYYYMMDD.xlsx로 저장. 별도 출석부 파일 불필요.
        """
        try:
            import re as _re
            wb = load_workbook(source_path, data_only=True)

            sheet_name = None
            for name in wb.sheetnames:
                if '출석부' in name:
                    sheet_name = name
                    break

            if not sheet_name:
                wb.close()
                return {'success': False, 'message': "'출석부양식' 시트를 찾을 수 없습니다.", 'count': 0}

            ws = wb[sheet_name]

            # 날짜 추출
            if not date_str:
                match = _re.search(r'(\d{6,8})', os.path.basename(source_path))
                date_str = match.group(1) if match else datetime.now().strftime('%Y%m%d')

            # 데이터 행 수 (이름이 있는 행까지)
            data_rows = 0
            for row_idx in range(2, ws.max_row + 1):
                name_val = ws.cell(row=row_idx, column=2).value
                if name_val and str(name_val).strip() and str(name_val) != 'NaN':
                    data_rows = row_idx
                else:
                    break

            # 새 워크북 생성, A~I열(1~9) / 헤더(1행) + 데이터(2~N+1행) 복사
            from openpyxl import Workbook
            new_wb = Workbook()
            new_ws = new_wb.active
            new_ws.title = "출석부양식"

            thin_border = Border(
                left=Side(style='thin'), right=Side(style='thin'),
                top=Side(style='thin'), bottom=Side(style='thin')
            )
            center_align = Alignment(horizontal='center', vertical='center')

            for row_idx in range(1, data_rows + 1):
                for col_idx in range(1, 10):  # A~I = 1~9
                    src_cell = ws.cell(row=row_idx, column=col_idx)
                    dst_cell = new_ws.cell(row=row_idx, column=col_idx,
                                           value=src_cell.value)
                    dst_cell.border = thin_border
                    dst_cell.alignment = center_align
                    if src_cell.font:
                        dst_cell.font = Font(
                            bold=src_cell.font.bold,
                            size=src_cell.font.size or 10)

            # 열 너비 조정
            col_widths = {'A': 5, 'B': 10, 'C': 12, 'D': 14, 'E': 6,
                          'F': 8, 'G': 10, 'H': 10, 'I': 8}
            for col_letter, width in col_widths.items():
                new_ws.column_dimensions[col_letter].width = width

            wb.close()

            # 저장
            if not output_dir:
                output_dir = os.path.dirname(source_path)
            os.makedirs(output_dir, exist_ok=True)

            person_count = data_rows - 1  # 헤더 제외
            filename = f"출석부양식_{date_str}.xlsx"
            save_path = os.path.join(output_dir, filename)
            new_wb.save(save_path)
            new_wb.close()

            return {
                'success': True,
                'message': f'출석부 저장 완료: {filename} ({person_count}명)',
                'count': person_count,
                'saved_path': save_path
            }

        except Exception as e:
            return {'success': False, 'message': str(e), 'count': 0}

    def get_summary(self, source_path: str) -> dict:
        """종합 통계 조회"""
        data = self.read_source_data(source_path)
        
        summary = {
            'total': len(data),
            'normal': 0,      # 일반 (본인부담)
            'vulnerable': 0,  # 취약계층
            'by_type': {
                '55세': 0,
                '장기실업': 0,
                '기초생활': 0,
                '장애인': 0,
                '20세': 0
            },
            'by_payment': {
                '현금': 0,
                '카드': 0,
                '계좌': 0
            }
        }
        
        for row in data:
            # 취약계층 분류
            vul = str(row.get('vulnerable', '0'))
            if vul in ['0', '', 'NaN', 'None']:
                summary['normal'] += 1
            else:
                summary['vulnerable'] += 1
                if '55' in vul:
                    summary['by_type']['55세'] += 1
                elif '장기' in vul or '실업' in vul:
                    summary['by_type']['장기실업'] += 1
                elif '기초' in vul:
                    summary['by_type']['기초생활'] += 1
                elif '장애' in vul:
                    summary['by_type']['장애인'] += 1
                elif '20' in vul:
                    summary['by_type']['20세'] += 1
            
            # 결제 방식 분류
            payment = str(row.get('payment', ''))
            if '현금' in payment:
                summary['by_payment']['현금'] += 1
            elif '카드' in payment:
                summary['by_payment']['카드'] += 1
            elif '계좌' in payment:
                summary['by_payment']['계좌'] += 1
        
        return summary
